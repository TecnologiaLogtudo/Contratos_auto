from __future__ import annotations

import tempfile
from pathlib import Path
import openpyxl
import pytest

from backend.app.engine.excel_processor import ExcelProcessor
from backend.app.domain.models import ItemStatus


def test_parse_raw_workbook_mixed_rows():
    with tempfile.TemporaryDirectory() as tmpdir:
        file_path = Path(tmpdir) / "teste_entrada.xlsx"

        # Cria uma planilha no formato bruto (linhas 1 a 7 cabeçalho legado, linha 8 em diante dados)
        wb = openpyxl.Workbook()
        ws = wb.active

        # Preenche linhas decorativas 1 a 6 e cabeçalho na linha 7
        for r in range(1, 7):
            ws.append([f"Cabecalho_{r}"])

        # Linha 7: Cabeçalho BSoft real
        header_row = [
            "Status", "Nro Cotação", "", "", "", "", "Validade", "", "", "",
            "Categoria Veículo", "Remetente", "Cidade / UF", "", "", "",
            "Frete a Pagar", "", "Frete Negociado", "", "", "",
            "Motorista / Placa"
        ]
        ws.append(header_row)

        # Linha 8: Item Válido Lactalis
        row_valid_lactalis = [
            None, "1832067", None, None, None, None, "15/08/2026", None, None, None,
            "TOCO", "LACTALIS DO BRASIL", "SALVADOR / BA", None, None, None,
            "R$ 500,00", None, "R$ 600,00", None, None, None,
            "JOAO PEDRO BRA2E19 15/08/2026"
        ]
        ws.append(row_valid_lactalis)

        # Linha 9: Item com Placa Faltante (deve ir para validação com erro)
        row_missing_placa = [
            None, "1832068", None, None, None, None, "15/08/2026", None, None, None,
            "TOCO", "LACTALIS DO BRASIL", "SALVADOR / BA", None, None, None,
            "R$ 500,00", None, "R$ 600,00", None, None, None,
            "MARCOS SEM PLACA 15/08/2026"
        ]
        ws.append(row_missing_placa)

        # Linha 10: Item Latam Válido (sem validade, mas com data pagamento)
        row_valid_latam = [
            None, "1832069", None, None, None, None, None, None, None, None,
            "CARRETA", "TAM LINHAS AEREAS", "JOAO PESSOA / PB", None, None, None,
            "R$ 1000,00", None, "R$ 1100,00", None, None, None,
            "CARLOS PEREIRA XYZ-9876 20/08/2026"
        ]
        ws.append(row_valid_latam)

        # Linha 11: Item com cotação vazia (deve ser ignorado)
        row_empty_quote = [None, "", None, None, None, None, None]
        ws.append(row_empty_quote)

        wb.save(file_path)
        wb.close()

        processor = ExcelProcessor(file_path)
        items = processor.load_and_parse()

        # Itens válidos carregados
        assert len(items) == 2
        assert items[0].nro_cotacao == "1832067"
        assert items[0].placa == "BRA2E19"
        assert items[0].cidade == "Salvador"
        assert items[0].status == ItemStatus.PENDENTE

        assert items[1].nro_cotacao == "1832069"
        assert items[1].placa == "XYZ9876"
        assert items[1].cidade == "J. Pessoa"

        # Item com erro de validação
        assert len(processor.validation_errors) == 1
        err_item, err_motivo = processor.validation_errors[0]
        assert err_item.nro_cotacao == "1832068"
        assert "Placa" in err_motivo

        # Testa salvamento
        out_path = Path(tmpdir) / "saida_processada.xlsx"
        saved = processor.save_output(out_path)
        assert saved.exists()

        wb_res = openpyxl.load_workbook(saved)
        ws_err = wb_res["Contrato não realizado"]
        err_rows = list(ws_err.iter_rows(values_only=True))
        assert any(r[0] == "1832068" and "Placa" in r[7] for r in err_rows[1:])
        wb_res.close()


def test_parse_real_xls_file_if_exists():
    real_file = Path(__file__).resolve().parents[1] / "planilha_pernoite - Copia.xls"
    if not real_file.exists():
        pytest.skip("Planilha de exemplo real não encontrada no diretório")

    processor = ExcelProcessor(real_file)
    items = processor.load_and_parse()
    assert len(items) > 0

    first_item = items[0]
    assert first_item.nro_cotacao
    assert first_item.placa
    assert first_item.status == ItemStatus.PENDENTE

