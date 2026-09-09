from __future__ import annotations

import tempfile
from pathlib import Path
import openpyxl
import pytest

from backend.app.engine.excel_processor import (
    processar_cidade_uf,
    processar_nome_placa,
    padronizar_cidade,
    ExcelProcessor,
)
from backend.app.domain.models import ItemContrato, ItemStatus


def test_processar_placas():
    # Placa Antiga
    nome, placa, dt, extra = processar_nome_placa("JOAO DA SILVA ABC-1234 15/08/2026")
    assert placa == "ABC1234"
    assert "JOAO DA SILVA" in nome
    assert dt == "15/08/2026"
    assert extra == "Não"

    # Placa Mercosul
    nome, placa, dt, extra = processar_nome_placa("MARIA SOUZA BRA2E19 10/09/2026 VIAGEM EXTRA")
    assert placa == "BRA2E19"
    assert "MARIA SOUZA" in nome
    assert dt == "10/09/2026"
    assert extra == "Sim"

    # Placa com espaço
    nome, placa, dt, extra = processar_nome_placa("CARLOS PEREIRA XYZ 9876")
    assert placa == "XYZ9876"
    assert "CARLOS PEREIRA" in nome


def test_limpeza_nome_motorista():
    # Remove pernoite, NF, horários
    raw = "pernoite 53640535 nf-115452 FELIPE LEAL FERREIRA JRK-1055 14h30"
    nome, placa, dt, extra = processar_nome_placa(raw)
    assert placa == "JRK1055"
    assert "FELIPE LEAL FERREIRA" in nome
    assert "pernoite" not in nome.lower()
    assert "nf" not in nome.lower()


def test_padronizar_cidades():
    assert padronizar_cidade("JOAO PESSOA") == "J. Pessoa"
    assert padronizar_cidade("João Pessoa") == "J. Pessoa"
    assert padronizar_cidade("Bayeux") == "J. Pessoa"
    assert padronizar_cidade("CAMACARI") == "Salvador"
    assert padronizar_cidade("Camaçari") == "Salvador"
    assert padronizar_cidade("SIMOES FILHO") == "Salvador"
    assert padronizar_cidade("Simões Filho") == "Salvador"
    assert padronizar_cidade("SERRA") == "Vitória"
    assert padronizar_cidade("Vila Velha") == "Vitória"
    assert padronizar_cidade("Feira de Santana") == "Feira de Santana"


def test_processar_cidade_uf():
    cid, uf = processar_cidade_uf("CAMACARI / BA")
    assert cid == "Salvador"
    assert uf == "BA"

    cid, uf = processar_cidade_uf("JOAO PESSOA / PB")
    assert cid == "J. Pessoa"
    assert uf == "PB"


def test_excel_processor_save_output():
    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = Path(tmpdir) / "resultado_test.xlsx"
        dummy_in = Path(tmpdir) / "dummy.xlsx"

        # Cria planilha dummy
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dados Processados"
        wb.save(dummy_in)
        wb.close()

        processor = ExcelProcessor(dummy_in)
        item1 = ItemContrato(
            nro_cotacao="1001",
            nome="MOTORISTA UM",
            placa="ABC1234",
            cidade="Salvador",
            uf="BA",
            remetente="Lactalis",
            status=ItemStatus.CONCLUIDO,
        )
        item2 = ItemContrato(
            nro_cotacao="1002",
            nome="MOTORISTA DOIS",
            placa="XYZ9876",
            cidade="Vitória",
            uf="ES",
            remetente="Latam",
            status=ItemStatus.ERRO,
            observacao_erro="Cotação bloqueada",
        )

        processor.items = [item1, item2]
        saved_file = processor.save_output(out_path)

        assert saved_file.exists()

        # Verifica conteúdo salvo
        wb_check = openpyxl.load_workbook(saved_file)
        assert "Dados Processados" in wb_check.sheetnames
        assert "Contrato não realizado" in wb_check.sheetnames

        ws_proc = wb_check["Dados Processados"]
        ws_err = wb_check["Contrato não realizado"]

        # Item 1 deve estar em Dados Processados
        rows_proc = list(ws_proc.iter_rows(values_only=True))
        assert any(r[0] == "1001" for r in rows_proc[1:])

        # Item 2 deve estar em Contrato não realizado com o motivo
        rows_err = list(ws_err.iter_rows(values_only=True))
        assert any(r[0] == "1002" and "Cotação bloqueada" in r[7] for r in rows_err[1:])

        wb_check.close()
