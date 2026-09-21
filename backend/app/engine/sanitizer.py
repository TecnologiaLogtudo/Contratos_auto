"""Módulo de auto-detecção e sanitização de planilhas de relatórios BSoft/LogTudo (.xls e .xlsx).

Garante que planilhas com layouts variantes (cabeçalhos deslocados, células mescladas,
rodapés, colunas ausentes ou em ordens diferentes) sejam normalizadas de forma determinística
antes da extração dos contratos.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, List, Optional, Sequence, Tuple
import openpyxl
import xlrd

TOTAL_MARKERS = ("totais", "total", "subtotal")
FOOTER_MARKERS = ("página", "pagina", "e-login", "impresso")

# Palavras-chave para auto-detecção da linha de cabeçalho
HEADER_HINTS = (
    "nro da cota", "cotação", "cotacao", "validade", "categoria",
    "remetente", "destinatário", "destinatario", "frete", "cliente",
    "nota fiscal", "nf", "veículo", "veiculo", "motorista", "placa",
    "cidade", "uf", "dados",
)

# Mapeamento canônico de sinônimos de colunas para identificação semântica
COLUMN_SYNONYMS = {
    "status": ["status", "situação", "situacao", "estado", "fase", "posição", "posicao"],
    "nro_cotacao": ["nro cotação", "nro cotacao", "nro da cotação", "nro da cotacao", "cotação", "cotacao", "ravex", "nro pedido"],
    "categoria_veiculo": ["categoria veículo", "categoria veiculo", "categoria", "tipo veículo", "tipo veiculo"],
    "cidade_uf": ["cidade/uf", "cidade / uf", "cidade", "municipio", "origem/destino"],
    "nome_placa": ["motorista/placa", "nome/placa", "motorista / placa", "nome / placa", "motorista", "placa"],
    "remetente": ["remetente", "cliente", "empresa", "cnpj remetente"],
    "validade": ["validade", "dt validade", "data validade", "data de validade"],
    "data_pagamento": ["data pagamento", "dt pagamento", "data previsto pagto", "previsao pagto"],
    "frete_a_pagar": ["frete a pagar", "frete pagar", "valor a pagar", "a pagar"],
    "frete_negociado": ["frete negociado", "valor negociado", "negociado"],
    "viagem_extra": ["viagem extra", "extra"],
    "observacao_interna": ["observação interna", "observacao interna", "obs interna", "observação", "observacao"],
}


@dataclass
class NormalizedSpreadsheet:
    headers: List[str]
    rows: List[List[Any]]
    header_row_index: int
    column_mapping: dict[str, int] = field(default_factory=dict)
    warnings: List[str] = field(default_factory=list)


def detect_header_row_xls(sheet: xlrd.sheet.Sheet, scan_limit: int = 20) -> int:
    """Detecta automaticamente o índice da linha de cabeçalho em sheets xlrd."""
    best, best_score = 0, 0
    for r in range(min(scan_limit, sheet.nrows)):
        score = 0
        for c in range(sheet.ncols):
            v = sheet.cell_value(r, c)
            if isinstance(v, str):
                lv = v.strip().lower()
                if any(h in lv for h in HEADER_HINTS):
                    score += 1
        if score > best_score:
            best, best_score = r, score
    return best if best_score >= 2 else 0


def detect_header_row_openpyxl(ws: openpyxl.worksheet.worksheet.Worksheet, scan_limit: int = 20) -> int:
    """Detecta automaticamente o índice (1-based) da linha de cabeçalho em openpyxl."""
    best, best_score = 1, 0
    for r_idx, row in enumerate(ws.iter_rows(max_row=scan_limit, values_only=True), start=1):
        score = 0
        for cell in row:
            if isinstance(cell, str):
                lv = cell.strip().lower()
                if any(h in lv for h in HEADER_HINTS):
                    score += 1
        if score > best_score:
            best, best_score = r_idx, score
    return best if best_score >= 2 else 1


def map_columns(headers: Sequence[str]) -> dict[str, int]:
    """Mapeia os campos canônicos do sistema para os índices de coluna encontrados."""
    mapping = {}
    normalized_headers = [re.sub(r'\s+', ' ', str(h or "").strip().lower()) for h in headers]

    for canonical_name, synonyms in COLUMN_SYNONYMS.items():
        for idx, h in enumerate(normalized_headers):
            if any(syn in h for syn in synonyms):
                mapping[canonical_name] = idx
                break

    return mapping


def sanitize_bsoft_xls(filepath: Path | str) -> NormalizedSpreadsheet:
    """Sanitiza e normaliza arquivo .xls desfazendo mesclagens e detectando cabeçalhos."""
    src = Path(filepath)
    wb = xlrd.open_workbook(str(src), formatting_info=True)
    sheet = wb.sheet_by_index(0)

    header_row = detect_header_row_xls(sheet)
    
    # Mapa de células mescladas
    merged_map = {}
    for (r0, r1, c0, c1) in sheet.merged_cells:
        for r in range(r0, r1):
            for c in range(c0, c1):
                merged_map[(r, c)] = (r0, r1, c0, c1)

    def get_val(r: int, c: int) -> Any:
        m = merged_map.get((r, c))
        if m is None:
            return sheet.cell_value(r, c)
        return sheet.cell_value(m[0], m[2])

    raw_headers = [get_val(header_row, c) for c in range(sheet.ncols)]
    headers = [str(h).replace("\n", " ").strip() if h else f"col_{i}" for i, h in enumerate(raw_headers)]

    rows: List[List[Any]] = []
    footer_hit = False

    for r in range(header_row + 1, sheet.nrows):
        cells = [get_val(r, c) for c in range(sheet.ncols)]

        # Ignora rodapé
        if not footer_hit and any(isinstance(v, str) and any(mk in v.lower() for mk in FOOTER_MARKERS) for v in cells):
            footer_hit = True
        if footer_hit:
            continue

        # Ignora totais
        first_text = next((str(v).strip().lower() for v in cells if isinstance(v, str) and v.strip()), "")
        if first_text and any(first_text.startswith(mk) for mk in TOTAL_MARKERS):
            continue

        # Ignora linhas totalmente vazias
        if all(v in ("", None) for v in cells):
            continue

        clean_row = []
        for c in range(sheet.ncols):
            v = cells[c]
            bs = merged_map.get((r, c))
            if bs is not None and (r, c) != (bs[0], bs[2]):
                anchor = cells[bs[2]]
                if v == anchor:
                    v = ""
            if isinstance(v, str):
                v = re.sub(r"\s+", " ", v.replace("\t", " ")).strip()
            clean_row.append(v)

        rows.append(clean_row)

    col_map = map_columns(headers)
    return NormalizedSpreadsheet(
        headers=headers,
        rows=rows,
        header_row_index=header_row,
        column_mapping=col_map,
    )


def sanitize_bsoft_xlsx(filepath: Path | str) -> NormalizedSpreadsheet:
    """Sanitiza e normaliza arquivo .xlsx detectando cabeçalhos e limpando ruídos."""
    src = Path(filepath)
    wb = openpyxl.load_workbook(src, data_only=True)
    try:
        ws = wb.active
        header_row_1based = detect_header_row_openpyxl(ws)

        headers: List[str] = []
        for idx, cell in enumerate(ws[header_row_1based], start=1):
            val = str(cell.value or "").replace("\n", " ").strip()
            headers.append(val if val else f"col_{idx}")

        rows: List[List[Any]] = []
        footer_hit = False

        for row in ws.iter_rows(min_row=header_row_1based + 1, values_only=True):
            if not row or all(c in (None, "") for c in row):
                continue

            if not footer_hit and any(isinstance(v, str) and any(mk in v.lower() for mk in FOOTER_MARKERS) for v in row):
                footer_hit = True
            if footer_hit:
                continue

            first_text = next((str(v).strip().lower() for v in row if isinstance(v, str) and v.strip()), "")
            if first_text and any(first_text.startswith(mk) for mk in TOTAL_MARKERS):
                continue

            clean_row = [str(c).strip() if isinstance(c, str) else c for c in row]
            rows.append(clean_row)

        col_map = map_columns(headers)
        return NormalizedSpreadsheet(
            headers=headers,
            rows=rows,
            header_row_index=header_row_1based - 1,
            column_mapping=col_map,
        )
    finally:
        wb.close()
